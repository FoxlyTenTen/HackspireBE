import io

from openpyxl import load_workbook
from sqlalchemy import delete as sa_delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ApiError
from app.db.models import Profile, TimelineItem
from app.schemas.admin import (
    TimelineImportConfirmResponse,
    TimelineImportConfirmRowIn,
    TimelineImportPreviewResponse,
    TimelineImportRow,
)
from app.services.audit import log_admin_action

REQUIRED_COLUMNS = ("bootcamp_name", "date_label", "time_session", "session_type")
ALL_COLUMNS = REQUIRED_COLUMNS + (
    "speaker_name",
    "speaker_bio",
    "organization",
    "session_description",
    "speaker_photo_url",
    "status",
    "remarks",
    "order_index",
)


def _parse_workbook(content: bytes) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ApiError(400, "File could not be parsed as an Excel (.xlsx) workbook")
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise ApiError(400, "The uploaded sheet is empty")
    columns = {
        str(cell).strip().lower(): idx
        for idx, cell in enumerate(header)
        if cell is not None
    }
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise ApiError(
            400,
            f"Missing required column(s) in header row: {', '.join(missing)}. "
            f"Expected columns: {', '.join(REQUIRED_COLUMNS)}",
        )

    def cell(row: tuple, key: str) -> str | None:
        idx = columns.get(key)
        if idx is None or idx >= len(row):
            return None
        value = row[idx]
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    parsed = []
    for row_number, row in enumerate(rows, start=2):
        if all(v is None for v in row):
            continue
        data = {key: cell(row, key) for key in ALL_COLUMNS}
        data["row_number"] = row_number
        parsed.append(data)
    if not parsed:
        raise ApiError(400, "The uploaded sheet has no data rows")
    return parsed


def _validate_rows(rows: list[dict]) -> list[TimelineImportRow]:
    validated: list[TimelineImportRow] = []
    for idx, row in enumerate(rows):
        issues: list[str] = []
        for field in REQUIRED_COLUMNS:
            if not row.get(field):
                issues.append(f"missing {field}")
        order_index_raw = row.get("order_index")
        try:
            order_index = int(order_index_raw) if order_index_raw else idx
        except ValueError:
            issues.append("order_index must be a number")
            order_index = idx
        status = row.get("status") or "TBD"
        if status not in ("TBD", "Proposed", "Confirmed", "Declined"):
            issues.append("status must be one of TBD/Proposed/Confirmed/Declined")
        validated.append(
            TimelineImportRow(
                row_number=row["row_number"],
                bootcamp_name=row.get("bootcamp_name"),
                date_label=row.get("date_label"),
                time_session=row.get("time_session"),
                session_type=row.get("session_type"),
                speaker_name=row.get("speaker_name"),
                speaker_bio=row.get("speaker_bio"),
                organization=row.get("organization"),
                session_description=row.get("session_description"),
                speaker_photo_url=row.get("speaker_photo_url"),
                status=status,
                remarks=row.get("remarks"),
                order_index=order_index,
                row_status="flagged" if issues else "ok",
                issues=issues,
            )
        )
    return validated


def preview(content: bytes) -> TimelineImportPreviewResponse:
    """Step 1: parse and validate WITHOUT writing anything."""
    rows = _parse_workbook(content)
    validated = _validate_rows(rows)
    ok = sum(1 for r in validated if r.row_status == "ok")
    return TimelineImportPreviewResponse(
        rows=validated, ok_count=ok, flagged_count=len(validated) - ok
    )


async def confirm(
    db: AsyncSession, admin: Profile, rows: list[TimelineImportConfirmRowIn]
) -> TimelineImportConfirmResponse:
    """Step 2: replace the entire timeline with the confirmed rows."""
    if not rows:
        raise ApiError(400, "No rows provided to import")
    await db.execute(sa_delete(TimelineItem))
    new_items = [
        TimelineItem(
            bootcamp_name=r.bootcamp_name,
            date_label=r.date_label,
            time_session=r.time_session,
            session_type=r.session_type,
            speaker_name=r.speaker_name,
            speaker_bio=r.speaker_bio,
            organization=r.organization,
            session_description=r.session_description,
            speaker_photo_url=r.speaker_photo_url,
            status=r.status,
            remarks=r.remarks,
            order_index=r.order_index,
        )
        for r in rows
    ]
    db.add_all(new_items)
    log_admin_action(
        db,
        admin.id,
        "timeline_imported",
        target_table="timeline_items",
        details=f"{len(new_items)} item(s)",
    )
    await db.commit()
    return TimelineImportConfirmResponse(created=len(new_items))
