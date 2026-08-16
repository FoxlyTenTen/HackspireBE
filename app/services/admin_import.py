import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ApiError
from app.db.models import Profile
from app.schemas.admin import (
    ImportConfirmRowIn,
    ImportConfirmResponse,
    ImportPreviewResponse,
    ImportRow,
)
from app.services import supabase_auth
from app.services.audit import log_admin_action

REQUIRED_COLUMNS = ("name", "email", "student_staff_id")
OPTIONAL_COLUMNS = ("reference_id", "university_name", "technical_skills", "preferred_role")


def _parse_workbook(content: bytes) -> list[dict]:
    """Read the first sheet; header row must contain name, email, student_staff_id."""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ApiError(400, "File could not be parsed as an Excel (.xlsx) workbook")
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise ApiError(400, "The uploaded sheet is empty")
    columns = {
        str(cell).strip().lower(): idx
        for idx, cell in enumerate(header)
        if cell is not None
    }
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise ApiError(
            400,
            f"Missing required column(s) in header row: {', '.join(missing)}. "
            f"Expected columns: {', '.join(REQUIRED_COLUMNS)}",
        )

    def cell(row: tuple, key: str) -> str | None:
        idx = columns.get(key)
        if idx is None or idx >= len(row):
            return None
        value = row[idx]
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    parsed = []
    for row_number, row in enumerate(rows, start=2):
        if all(v is None for v in row):
            continue  # skip fully blank rows
        parsed.append(
            {
                "row_number": row_number,
                "name": cell(row, "name"),
                "email": cell(row, "email"),
                "student_staff_id": cell(row, "student_staff_id"),
                "reference_id": cell(row, "reference_id"),
                "university_name": cell(row, "university_name"),
                "technical_skills": cell(row, "technical_skills"),
                "preferred_role": cell(row, "preferred_role"),
            }
        )
    if not parsed:
        raise ApiError(400, "The uploaded sheet has no data rows")
    return parsed


async def _existing_sets(db: AsyncSession) -> tuple[set[str], set[str]]:
    emails = {e.lower() for e in (await db.scalars(select(Profile.email))).all()}
    ids = {s.lower() for s in (await db.scalars(select(Profile.student_staff_id))).all()}
    return emails, ids


def _validate_rows(
    rows: list[dict], db_emails: set[str], db_ids: set[str]
) -> list[ImportRow]:
    """Flag each row individually — a bad row never rejects the whole file."""
    seen_emails: set[str] = set()
    seen_ids: set[str] = set()
    validated: list[ImportRow] = []
    for row in rows:
        issues: list[str] = []
        for field in REQUIRED_COLUMNS:
            if not row[field]:
                issues.append(f"missing {field}")
        email = (row["email"] or "").lower()
        sid = (row["student_staff_id"] or "").lower()
        if email:
            if email in db_emails:
                issues.append("email already registered")
            if email in seen_emails:
                issues.append("duplicate email within file")
            seen_emails.add(email)
        if sid:
            if sid in db_ids:
                issues.append("student_staff_id already registered")
            if sid in seen_ids:
                issues.append("duplicate student_staff_id within file")
            seen_ids.add(sid)
        validated.append(
            ImportRow(
                row_number=row["row_number"],
                name=row["name"],
                email=row["email"],
                student_staff_id=row["student_staff_id"],
                reference_id=row.get("reference_id"),
                university_name=row.get("university_name"),
                technical_skills=row.get("technical_skills"),
                preferred_role=row.get("preferred_role"),
                status="flagged" if issues else "ok",
                issues=issues,
            )
        )
    return validated


async def preview(db: AsyncSession, content: bytes) -> ImportPreviewResponse:
    """Step 1: parse and validate WITHOUT writing anything."""
    rows = _parse_workbook(content)
    db_emails, db_ids = await _existing_sets(db)
    validated = _validate_rows(rows, db_emails, db_ids)
    ok = sum(1 for r in validated if r.status == "ok")
    return ImportPreviewResponse(
        rows=validated, ok_count=ok, flagged_count=len(validated) - ok
    )


async def confirm(
    db: AsyncSession, admin: Profile, rows: list[ImportConfirmRowIn]
) -> ImportConfirmResponse:
    """Step 2: re-validate clean rows, then invite each via Supabase Auth.

    invite_user_by_email sends a real invite email; the participant sets their
    own password from its link, which replaces the old dummy-password +
    forced-first-login flow. Rows whose invite call fails are NOT created (a
    profile without an auth user could never log in) — they come back in
    `skipped_rows` with the failure reason so the admin can retry them.
    """
    if not rows:
        raise ApiError(400, "No rows provided to import")
    raw = [
        {
            "row_number": i,
            "name": r.name.strip() or None,
            "email": r.email.strip() or None,
            "student_staff_id": r.student_staff_id.strip() or None,
            "reference_id": (r.reference_id or "").strip() or None,
            "university_name": (r.university_name or "").strip() or None,
            "technical_skills": (r.technical_skills or "").strip() or None,
            "preferred_role": (r.preferred_role or "").strip() or None,
        }
        for i, r in enumerate(rows, start=1)
    ]
    db_emails, db_ids = await _existing_sets(db)
    validated = _validate_rows(raw, db_emails, db_ids)

    created = 0
    invite_failed = 0
    skipped_rows: list[ImportRow] = []
    for row in validated:
        if row.status != "ok":
            skipped_rows.append(row)
            continue
        try:
            auth_user = await supabase_auth.invite_user_by_email(
                row.email,
                data={"name": row.name, "student_staff_id": row.student_staff_id},
            )
        except ApiError as exc:
            invite_failed += 1
            row.status = "flagged"
            row.issues = [f"invite email failed: {exc.message}"]
            skipped_rows.append(row)
            continue
        db.add(
            Profile(
                id=uuid.UUID(auth_user["id"]),  # profiles.id matches auth.users.id
                name=row.name,
                email=row.email,
                student_staff_id=row.student_staff_id,
                reference_id=row.reference_id,
                university_name=row.university_name,
                technical_skills=row.technical_skills,
                preferred_role=row.preferred_role,
                role="participant",
                first_login=False,  # they set their own password via the invite link
                invite_status="invited",
            )
        )
        created += 1
    log_admin_action(
        db,
        admin.id,
        "participant_imported",
        target_table="profiles",
        details=(
            f"{created} invited, {len(skipped_rows)} skipped"
            f" ({invite_failed} invite failures)"
        ),
    )
    await db.commit()
    return ImportConfirmResponse(
        created=created,
        skipped=len(skipped_rows),
        invite_failed=invite_failed,
        skipped_rows=skipped_rows,
    )
